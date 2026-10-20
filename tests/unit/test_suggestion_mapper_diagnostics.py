"""Synthetic fixtures verify diagnostics, not any historical LLM incident."""

import copy
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from celery.app.log import TaskFormatter
from celery.utils.log import ColorFormatter

from app.core.logging_setup import CorrelationJsonFormatter, PlainCorrelationFormatter
from app.infrastructure.db.models.enums import ChangeType, SuggestionStatus
from app.workers.pipeline import suggestion_diagnostics as diagnostics
from app.workers.pipeline import suggestion_mapper as mapper
from app.workers.celery_app import celery_app


import logging
import uuid

from app.domain.interfaces.llm_client import LLMSuggestionBatch, LLMSuggestionItem
from app.workers.pipeline.suggestion_mapper import map_to_suggestions

LOGGER = "syncscribe.workers.suggestion_mapper"
JOB = uuid.UUID("11111111-1111-1111-1111-111111111111")
DOCUMENT_ID = uuid.UUID("22222222-2222-2222-2222-222222222222")
SOURCE = uuid.UUID("33333333-3333-3333-3333-333333333333")


def item(kind="modify", old="missing", new="replacement", **kwargs):
    return LLMSuggestionItem(
        section_ref=kwargs.pop("section_ref", "absent section"),
        change_type=kind,
        old_text=old,
        new_text=new,
        **kwargs,
    )


def test_skipped_item_has_zero_based_index_in_console_message(caplog):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        result = map_to_suggestions(
            LLMSuggestionBatch(items=[item(old="present"), item()]),
            JOB,
            DOCUMENT_ID,
            document_text="present",
        )
    assert len(result) == 1
    assert any("item_index_0based=1" in record.getMessage() for record in caplog.records)


def run(*items, document="present", document_format="txt", enabled=True):
    return map_to_suggestions(
        LLMSuggestionBatch(items=list(items), raw_response="PRIVATE_RAW_RESPONSE"),
        JOB,
        DOCUMENT_ID,
        source_id=SOURCE,
        document_text=document,
        document_format=document_format,
        diagnostics_enabled=enabled,
    )


def detail_records(caplog):
    return [r for r in caplog.records if hasattr(r, "item_index_0based")]


@pytest.mark.parametrize(
    "quote,document,category",
    [
        ("left\u00a0right", "left right", "space_variants"),
        ("left right", "left\u00a0right", "space_variants"),
        ("one\r\ntwo", "one\ntwo", "line_endings"),
        ("one\ntwo", "one\r\ntwo", "line_endings"),
        ("cafe\u0301 menu", "caf\u00e9 menu", "unicode_nfc"),
        ("caf\u00e9 menu", "cafe\u0301 menu", "unicode_nfc"),
        ("left\tright", "left right", "tabs"),
        ("left right", "left\tright", "tabs"),
        ("pre\u00adfix", "prefix", "invisible_chars"),
        ("pre\u200bfix", "prefix", "invisible_chars"),
        ("pre\ufefffix", "prefix", "invisible_chars"),
        ("prefix", "pre\u200bfix", "invisible_chars"),
        ("PRESENT", "present", "case_only"),
        ("present!", "present.", "unknown"),
        ("a paraphrase", "different wording", "unknown"),
        ("&lt;document>sample", "<document>sample", "prompt_delimiter_escape"),
        ("&lt;/source>sample", "</source>sample", "prompt_delimiter_escape"),
        ("&lt;div>sample", "<div>sample", "unknown"),
        ("<document>not present", "irrelevant", "unknown"),
        ("LEFT right", "left right | left right", "ambiguous"),
        ("AAA", "aaaa", "ambiguous"),
        ("left\u00a0right", "left right | left right", "ambiguous"),
    ],
)
def test_classification_does_not_accept_transformed_matches(caplog, quote, document, category):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert run(item(old=quote), document=document) == []
    record = detail_records(caplog)[0]
    assert record.mismatch_category == category
    assert record.reason == "old_text_not_in_document"
    assert record.checked_old_text_len == len(quote.strip())
    assert record.raw_old_text_len == len(quote)
    assert record.document_len == len(document)
    assert not any(
        key in record.__dict__
        for key in ("position", "offset", "common_prefix_len", "text_char", "document_char")
    )


@pytest.mark.parametrize("kind", ["add", "modify", "delete"])
@pytest.mark.parametrize("document_format", ["docx", "doc", "txt", "md", "markdown"])
@pytest.mark.parametrize("quote", ["one\ntwo", "one\r\ntwo"])
def test_multiline_rules_are_unchanged(caplog, kind, document_format, quote):
    suggestion = item(kind, quote, None if kind == "delete" else "replacement")
    result = run(suggestion, document=quote, document_format=document_format)
    rejected = kind != "add" and document_format in {"docx", "doc"}
    assert len(result) == (0 if rejected else 1)
    if rejected:
        assert detail_records(caplog)[0].reason == "old_text_spans_paragraphs"
        assert detail_records(caplog)[0].mismatch_category == "not_applicable"
    else:
        assert result[0].original_text == quote


@pytest.mark.parametrize("kind", ["modify", "delete", "add"])
@pytest.mark.parametrize("quote", ["present", "  present  ", "\n present\r\n"])
def test_literal_and_outer_whitespace_preserve_original_fields(kind, quote):
    result = run(item(kind, quote, None if kind == "delete" else "replacement"))
    assert len(result) == 1
    assert result[0].original_text == (quote.strip() if kind == "add" else quote)
    assert result[0].status is SuggestionStatus.PENDING


@pytest.mark.parametrize(
    "raw,origin,final",
    [
        (item("add", "legacy", insert_after="missing anchor"), "insert_after", "add"),
        (item("add", "missing anchor"), "legacy_old_text", "add"),
        (item("add", None), "missing", "add"),
        (item("modify", "missing", "missing appended"), "normalized_modify_tail", "add"),
        (item("modify", "missing", ""), "not_applicable", "delete"),
    ],
)
def test_raw_type_index_and_anchor_provenance_survive_normalization(caplog, raw, origin, final):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        result = run(item(old="present"), raw)
    assert len(result) == 1
    record = detail_records(caplog)[0]
    assert record.item_index_0based == 1
    assert record.original_change_type == raw.change_type
    assert record.final_change_type == final
    assert record.anchor_origin == origin
    assert record.raw_old_text_len == len(raw.old_text or "")
    assert record.raw_insert_after_len == len(raw.insert_after or "")
    assert record.analysis_job_id == str(JOB)
    assert record.document_id == str(DOCUMENT_ID)
    assert record.source_id == str(SOURCE)
    if origin in {"insert_after", "legacy_old_text", "normalized_modify_tail"}:
        normalized = mapper._normalize(raw)
        assert record.anchor_len == len(normalized.insert_after or normalized.old_text or "")


def test_inferred_anchor_origin_is_not_lost():
    raw = item("add", None, section_ref="Heading")
    normalized = mapper._normalize(raw)
    resolved, inferred = mapper._resolve_anchor(normalized, "# Heading\nparagraph")
    assert inferred
    assert resolved.insert_after is None
    assert diagnostics.anchor_origin(raw, normalized, inferred=inferred) == "inferred_section"
    result = run(raw, document="# Heading\nparagraph")
    assert result[0].original_text == "paragraph"
    assert "Место вставки выбрано автоматически" in result[0].rationale


@pytest.mark.parametrize("kind", ["add", "modify", "delete"])
def test_repeated_exact_anchor_is_still_accepted_without_selecting_position(kind):
    result = run(item(kind, "present"), document="present / present")
    assert len(result) == 1
    assert result[0].original_text == "present"


@pytest.mark.parametrize("kept", [0, 1])
def test_partial_and_total_drop_aggregates(caplog, kept):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        result = run(*([item(old="present")] * kept), item(), item("delete", None, None))
    assert len(result) == kept
    summary = next(r for r in caplog.records if hasattr(r, "received"))
    assert (summary.received, summary.kept, summary.dropped) == (kept + 2, kept, 2)
    assert summary.diagnostic_records == 2
    assert f"received={kept + 2} kept={kept} dropped=2" in summary.getMessage()
    assert any("Все 2 правок" in r.getMessage() for r in caplog.records) == (kept == 0)


def test_empty_batch_is_not_a_loss(caplog):
    assert run() == []
    assert not caplog.records


def test_log_count_and_diagnostic_work_are_capped(caplog, monkeypatch):
    calls = []
    original = diagnostics.DiagnosticBudget.classify

    def spy(self, quote, document):
        calls.append(len(quote) + len(document))
        return original(self, quote, document)

    monkeypatch.setattr(diagnostics.DiagnosticBudget, "classify", spy)
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert run(*[item() for _ in range(250)]) == []
    assert len(calls) == diagnostics.MAX_RECORDS
    assert len(detail_records(caplog)) == diagnostics.MAX_RECORDS
    assert len(caplog.records) == diagnostics.MAX_RECORDS + 2
    summary = next(r for r in caplog.records if hasattr(r, "received"))
    assert summary.diagnostic_records_omitted == 230


@pytest.mark.parametrize("oversized", ["quote", "document"])
def test_input_limits_are_checked_before_any_transform(monkeypatch, oversized):
    def forbidden(value):
        raise AssertionError("transform must not run")

    monkeypatch.setattr(diagnostics, "_TRANSFORMS", (("test", forbidden),))
    quote = "x" * (diagnostics.MAX_QUOTE_CHARS + 1) if oversized == "quote" else "x"
    document = "z" * (diagnostics.MAX_DOCUMENT_CHARS + 1) if oversized == "document" else "z"
    assert diagnostics.DiagnosticBudget().classify(quote, document) == "input_limit"


def test_shared_budget_limits_total_transform_work(monkeypatch):
    calls = []

    def transform(value):
        calls.append(len(value))
        return value

    monkeypatch.setattr(diagnostics, "_TRANSFORMS", (("test", transform),))
    budget = diagnostics.DiagnosticBudget(remaining_chars=12)
    assert budget.classify("xxx", "zzz") == "unknown"
    assert budget.classify("xxx", "zzz") == "unknown"
    assert budget.classify("xxx", "zzz") == "budget_limit"
    assert sum(calls) == 12
    assert budget.remaining_chars == 0


def test_no_character_code_or_detailed_mode_by_default(caplog):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        run(item(old="pr\u200besent"))
    assert not hasattr(diagnostics, "first_divergence")
    assert not hasattr(mapper, "first_divergence")
    assert not hasattr(mapper, "_codepoint")
    assert not any("U+" in r.getMessage() for r in caplog.records)


def test_all_log_surfaces_are_metadata_only(caplog):
    markers = [
        "PRIVATE_DOCUMENT_CONTENT",
        "PRIVATE_QUOTE_CONTENT",
        "PRIVATE_NEW_CONTENT",
        "PRIVATE_ANCHOR_CONTENT",
        "PRIVATE_SECTION_CONTENT",
        "PRIVATE_REASON_CONTENT",
        "PRIVATE_RAW_RESPONSE",
        "PRIVATE_TYPE_CONTENT\nforged_log=true",
    ]
    raw = item(
        markers[-1],
        markers[1],
        markers[2],
        insert_after=markers[3],
        section_ref=markers[4],
        rationale=markers[5],
    )
    with caplog.at_level(logging.INFO, logger=LOGGER):
        assert (
            run(
                raw,
                item("add", markers[1], markers[2], insert_after=markers[3]),
                document=markers[0],
            )
            == []
        )
    assert detail_records(caplog)[0].original_change_type == "unknown"
    assert detail_records(caplog)[0].final_change_type == "unknown"
    for record in caplog.records:
        rendered = [
            record.getMessage(),
            repr(record.__dict__),
            CorrelationJsonFormatter().format(copy.copy(record)),
            PlainCorrelationFormatter().format(copy.copy(record)),
            TaskFormatter(celery_app.conf.worker_task_log_format).format(copy.copy(record)),
            ColorFormatter(celery_app.conf.worker_log_format).format(copy.copy(record)),
        ]
        assert record.exc_info is None
        assert record.exc_text is None
        for output in rendered:
            assert not any(marker in output for marker in markers)
            assert "forged_log" not in output
        if hasattr(record, "item_index_0based"):
            payload = json.loads(CorrelationJsonFormatter().format(copy.copy(record)))
            assert payload["item_index_0based"] == record.item_index_0based
            for formatter in (PlainCorrelationFormatter(), TaskFormatter("%(message)s")):
                assert "item_index_0based=" in formatter.format(copy.copy(record))
                assert f"source_id={SOURCE}" in formatter.format(copy.copy(record))


def test_invalid_ids_cannot_inject_logs(caplog):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert (
            map_to_suggestions(
                LLMSuggestionBatch(items=[item()]),
                "PRIVATE_JOB",
                "PRIVATE_DOCUMENT",
                source_id="PRIVATE_SOURCE",
                document_text="present",
            )
            == []
        )
    assert "PRIVATE" not in repr([r.__dict__ for r in caplog.records])
    assert detail_records(caplog)[0].source_id == "unavailable"


def columns(suggestions):
    return [
        {
            column.key: getattr(suggestion, column.key)
            for column in suggestion.__mapper__.column_attrs
        }
        for suggestion in suggestions
    ]


@pytest.mark.parametrize("document_format", ["docx", "txt", "md"])
def test_diagnostics_toggle_preserves_every_suggestion_column_and_status(document_format):
    batch = [
        item(old="present", rationale="kept rationale", confidence_score=0.7),
        item("delete", "  present\n", None),
        item("modify", "present", ""),
        item("modify", "present", "present appended"),
        item("add", None, section_ref="Heading"),
        item("add", "present", insert_after="one\ntwo"),
        item("modify", "one\ntwo", "replacement"),
        item(old="PRESENT"),
        item("add", None),
        item("unknown", "missing"),
    ]
    document = "# Heading\npresent\none\ntwo"
    enabled = run(*batch, document=document, document_format=document_format)
    disabled = run(*batch, document=document, document_format=document_format, enabled=False)
    assert columns(enabled) == columns(disabled)
    assert all(s.status is SuggestionStatus.PENDING for s in enabled)
    assert [s.change_type for s in enabled[:4]] == [
        ChangeType.MODIFY,
        ChangeType.DELETE,
        ChangeType.DELETE,
        ChangeType.ADD,
    ]
    assert batch[1].old_text == "  present\n"
    assert batch[3].change_type == "modify"
    assert all(not s.sources for s in enabled + disabled)


def test_disabled_diagnostics_do_not_classify_or_log_details(caplog, monkeypatch):
    def forbidden(*args):
        raise AssertionError("diagnostic path must not run")

    monkeypatch.setattr(diagnostics.DiagnosticBudget, "classify", forbidden)
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert run(item(), enabled=False) == []
    assert not detail_records(caplog)
    assert any("dropped=1" in r.getMessage() for r in caplog.records)


async def test_pipeline_forwards_source_id_without_real_llm_or_services(monkeypatch):
    from app.workers.tasks import analysis_tasks as tasks
    from app.infrastructure.db.models.enums import AnalysisJobStatus

    source = SimpleNamespace(id=SOURCE, type=None)
    document = SimpleNamespace(
        id=DOCUMENT_ID, storage_key="synthetic", format=SimpleNamespace(value="txt")
    )
    job = SimpleNamespace(document_id=DOCUMENT_ID, status=AnalysisJobStatus.PROCESSING)
    uow = SimpleNamespace(
        jobs=SimpleNamespace(
            get_by_id=AsyncMock(return_value=job),
            mark_processing_if_active=AsyncMock(return_value=True),
        ),
        sources=SimpleNamespace(get_by_id=AsyncMock(return_value=source)),
        documents=SimpleNamespace(get_by_id=AsyncMock(return_value=document)),
        job_sources=SimpleNamespace(
            ensure=AsyncMock(return_value=SimpleNamespace(id=uuid.uuid4())),
            get_completed_count=AsyncMock(return_value=None),
            claim=AsyncMock(return_value=uuid.uuid4()),
            mark_completed=AsyncMock(return_value=True),
        ),
        suggestions=SimpleNamespace(bulk_create=AsyncMock()),
        commit=AsyncMock(),
    )

    class Context:
        async def __aenter__(self):
            return uow

        async def __aexit__(self, *args):
            pass

    monkeypatch.setattr(tasks, "uow_factory", Context)
    monkeypatch.setattr(tasks, "_build_source_ref", lambda s: SimpleNamespace(type=None))
    monkeypatch.setattr(tasks, "_load_document_text", AsyncMock(return_value="present"))
    monkeypatch.setattr(
        tasks,
        "_get_connector_for",
        lambda t: SimpleNamespace(fetch=AsyncMock(return_value="synthetic source")),
    )
    llm = SimpleNamespace(
        generate_suggestions=AsyncMock(return_value=LLMSuggestionBatch(items=[item()]))
    )
    monkeypatch.setattr(tasks, "_get_llm_client", lambda: llm)
    forwarded = []

    def capture(*args, **kwargs):
        forwarded.append(kwargs["source_id"])
        return map_to_suggestions(*args, **kwargs)

    monkeypatch.setattr(tasks, "map_to_suggestions", capture)
    result = await tasks._process_source(str(JOB), str(SOURCE))
    assert result == {"source_id": str(SOURCE), "status": "ok", "count": 0}
    assert forwarded == [SOURCE]
    llm.generate_suggestions.assert_awaited_once_with("present", "synthetic source", "txt")
    uow.suggestions.bulk_create.assert_awaited_once_with([])


def test_classifier_failure_does_not_change_results_or_log_exception(caplog, monkeypatch):
    def fail(*args):
        raise ValueError("PRIVATE_DOCUMENT_AND_QUOTE")

    monkeypatch.setattr(diagnostics.DiagnosticBudget, "classify", fail)
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        result = run(item(old="present"), item())
    assert len(result) == 1
    assert detail_records(caplog)[0].mismatch_category == "diagnostic_error"
    assert "PRIVATE" not in repr([r.__dict__ for r in caplog.records])
    assert all(r.exc_info is None for r in caplog.records)


@pytest.mark.parametrize("quote", ["\ud800", "a\x00b", "\u200b", "\u00ad", "\t", ""])
def test_unusual_text_is_safe_metadata_only(caplog, quote):
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert run(item(old=quote)) == []
    for record in caplog.records:
        rendered = CorrelationJsonFormatter().format(copy.copy(record))
        rendered.encode("utf-8", errors="strict")
        assert record.exc_info is None


@pytest.mark.parametrize("kind", ["add", "modify", "delete"])
@pytest.mark.parametrize("document_format", ["docx", "txt", "md"])
def test_real_parser_multiline_contract_is_unchanged(kind, document_format):
    import io

    from docx import Document
    from app.infrastructure.parsers.docx_parser import DocxParser
    from app.infrastructure.parsers.txt_parser import TxtParser
    from app.infrastructure.parsers.markdown_parser import MarkdownParser

    if document_format == "docx":
        doc = Document()
        doc.add_paragraph("one")
        doc.add_paragraph("two")
        stream = io.BytesIO()
        doc.save(stream)
        parsed = DocxParser().parse(stream.getvalue())
    else:
        parser = TxtParser() if document_format == "txt" else MarkdownParser()
        parsed = parser.parse(b"one\r\ntwo")
    assert parsed.plain_text == "one\ntwo"
    result = run(
        item(kind, parsed.plain_text), document=parsed.plain_text, document_format=document_format
    )
    assert len(result) == (0 if document_format == "docx" and kind != "add" else 1)


def test_diagnostic_budget_is_shared_by_mapper_records(caplog, monkeypatch):
    monkeypatch.setattr(mapper, "DiagnosticBudget", lambda: diagnostics.DiagnosticBudget(0))
    with caplog.at_level(logging.WARNING, logger=LOGGER):
        assert run(item(), item()) == []
    assert [r.mismatch_category for r in detail_records(caplog)] == ["budget_limit"] * 2
